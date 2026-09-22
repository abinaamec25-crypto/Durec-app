"""
Microsoft Word (.docx) Document Parser.
Extracts records from embedded Word tables or formatted key-value paragraphs.
Uses python-docx.
"""
import io
import re
from typing import Any, Dict, List


class DOCXParser:
    """Parses .docx files, extracting records from tables or key-value sections."""

    def parse(self, file_path_or_bytes: Any, filename: str = "document.docx") -> List[Dict[str, Any]]:
        import docx

        if isinstance(file_path_or_bytes, bytes):
            doc = docx.Document(io.BytesIO(file_path_or_bytes))
        else:
            doc = docx.Document(file_path_or_bytes)

        records: List[Dict[str, Any]] = []

        # Strategy 1: Extract from document tables
        if doc.tables:
            for t_idx, table in enumerate(doc.tables):
                if len(table.rows) < 2:
                    continue
                header = [cell.text.strip() for cell in table.rows[0].cells]
                for r_idx, row in enumerate(table.rows[1:], start=2):
                    cell_vals = [cell.text.strip() for cell in row.cells]
                    if not any(cell_vals):
                        continue
                    item: Dict[str, Any] = {}
                    for c_idx, val in enumerate(cell_vals):
                        key = header[c_idx] if c_idx < len(header) and header[c_idx] else f"col_{c_idx}"
                        item[key] = val
                    item["__source_file__"] = filename
                    item["__file_type__"] = "Word (.docx)"
                    item["__line_number__"] = r_idx
                    records.append(item)

            if records:
                return records

        # Strategy 2: Extract from paragraphs (Key: Value blocks)
        current_record: Dict[str, Any] = {}
        kv_pattern = re.compile(r'^([A-Za-z0-9_\s]{2,25})\s*[:=\-]\s*(.+)$')

        for p_idx, p in enumerate(doc.paragraphs, start=1):
            text = p.text.strip()
            if not text:
                if current_record:
                    current_record["__source_file__"] = filename
                    current_record["__file_type__"] = "Word (.docx)"
                    current_record["__line_number__"] = p_idx
                    records.append(current_record)
                    current_record = {}
                continue

            m = kv_pattern.match(text)
            if m:
                k, v = m.group(1).strip(), m.group(2).strip()
                if k.lower() in [k_exist.lower() for k_exist in current_record.keys()]:
                    current_record["__source_file__"] = filename
                    current_record["__file_type__"] = "Word (.docx)"
                    current_record["__line_number__"] = p_idx
                    records.append(current_record)
                    current_record = {}
                current_record[k] = v

        if current_record:
            current_record["__source_file__"] = filename
            current_record["__file_type__"] = "Word (.docx)"
            current_record["__line_number__"] = len(doc.paragraphs)
            records.append(current_record)

        return records
