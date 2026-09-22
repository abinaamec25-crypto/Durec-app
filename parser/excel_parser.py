"""
Excel Spreadsheet Parser supporting modern .xlsx and legacy .xls files.
Uses openpyxl for .xlsx and xlrd for .xls.
"""
import io
import os
from typing import Any, Dict, List


class ExcelParser:
    """Parses Excel workbooks, extracting records from the first active worksheet."""

    def parse(self, file_path_or_bytes: Any, filename: str = "file.xlsx") -> List[Dict[str, Any]]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".xls":
            return self._parse_xls(file_path_or_bytes, filename)
        return self._parse_xlsx(file_path_or_bytes, filename)

    def _parse_xlsx(self, file_path_or_bytes: Any, filename: str) -> List[Dict[str, Any]]:
        import openpyxl

        if isinstance(file_path_or_bytes, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
        else:
            wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)

        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
        records: List[Dict[str, Any]] = []

        for row_idx, row in enumerate(rows[1:], start=2):
            if not any(c is not None for c in row):
                continue
            item: Dict[str, Any] = {}
            for col_idx, val in enumerate(row):
                key = header[col_idx] if col_idx < len(header) else f"col_{col_idx}"
                item[key] = "" if val is None else str(val).strip()

            item["__source_file__"] = filename
            item["__file_type__"] = "Excel (.xlsx)"
            item["__line_number__"] = row_idx
            records.append(item)

        return records

    def _parse_xls(self, file_path_or_bytes: Any, filename: str) -> List[Dict[str, Any]]:
        import xlrd

        if isinstance(file_path_or_bytes, bytes):
            wb = xlrd.open_workbook(file_contents=file_path_or_bytes)
        else:
            wb = xlrd.open_workbook(file_path_or_bytes)

        sheet = wb.sheet_by_index(0)
        if sheet.nrows == 0:
            return []

        header = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        records: List[Dict[str, Any]] = []

        for r in range(1, sheet.nrows):
            row_vals = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if not any(v != "" and v is not None for v in row_vals):
                continue
            item: Dict[str, Any] = {}
            for c, val in enumerate(row_vals):
                key = header[c] if c < len(header) else f"col_{c}"
                item[key] = "" if val is None else str(val).strip()

            item["__source_file__"] = filename
            item["__file_type__"] = "Excel (.xls)"
            item["__line_number__"] = r + 1
            records.append(item)

        return records
